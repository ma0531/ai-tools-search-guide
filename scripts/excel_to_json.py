# -*- coding: utf-8 -*-
"""
data/ai-tools.xlsx（Excel）を読み込んで、data/ai-tools.json を作り直すスクリプト。
GitHub Actions から自動で実行されます。

入力ミスがあると、どの行のどこが間違っているかを表示して止まります。
（JSONは書き換わらないので、サイトが壊れることはありません）
"""
import json
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "ai-tools.xlsx"
JSON_OUT = ROOT / "data" / "ai-tools.json"
SHEET = "AIツール一覧"

# サイトの「目的で絞り込む」にある項目と同じもの
# 新しい目的を増やすときは、ここと index.html の両方に追加してください
CATEGORIES = [
    "よく使われるAIのTOP10", "パーソナルアシスタント", "情報検索・リサーチ", "ノーコードでアプリ作成",
    "開発者向けツール", "コード補完", "Android開発", "ゲーム開発", "画像作成", "画像編集", "3Dモデル生成",
    "動画作成", "音楽・音声生成", "音声認識", "テキスト生成・AI執筆", "デザイン・UI/UX", "言語・翻訳",
    "要約と構造化", "データ分析・予測", "Excel・事務作業", "タスクの自動化（RPA連携）", "学習・教育",
    "悩み相談", "論文・リサーチ分析", "医療・ヘルスケア", "製造・ロボティクス", "日本語特化",
]
TOP10 = "よく使われるAIのTOP10"
PRICING = {"完全無料", "無料枠あり", "トライアルのみ"}
DEFAULT_LIMIT = {
    "完全無料": "無料で利用可能",
    "無料枠あり": "無料プランあり（回数・機能に制限）",
    "トライアルのみ": "無料トライアルのみ（継続利用は有料）",
}
HEADERS = {
    "AI名": "name", "URL": "url", "国": "country", "開発企業": "company", "カテゴリ": "category",
    "強み": "strengths", "料金タイプ": "pricing", "無料利用範囲": "freeLimit",
    "日本語対応": "japanese", "登録不要": "noSignup", "TOP10順位": "rank", "対応端末": "platform",
}
PLATFORM = {"PC＋スマホ": "both", "PCのみ": "pc", "スマホのみ": "mobile"}
# 英語版の列（なくても動きます）
OPTIONAL_HEADERS = {"AI名（英語）": "nameEn", "強み（英語）": "strengthsEn", "無料利用範囲（英語）": "freeLimitEn"}
REQUIRED = ["AI名", "URL", "国", "開発企業", "カテゴリ", "強み", "料金タイプ", "日本語対応", "対応端末"]


def text(v):
    if v is None:
        return ""
    return str(v).strip()


def split_list(s, seps):
    parts = re.split(seps, s)
    return [p.strip() for p in parts if p.strip()]


def normalize_japanese(s):
    s = s.replace("○", "◯").replace("〇", "◯").replace("O", "◯").replace("o", "◯")
    s = s.replace("✕", "×").replace("x", "×").replace("X", "×").replace("▲", "△")
    return s


def to_bool(s):
    return text(s) in {"◯", "○", "〇", "はい", "TRUE", "True", "true", "1", "Yes", "yes"}


def main():
    if not XLSX.exists():
        print(f"❌ {XLSX} が見つかりません")
        sys.exit(1)

    wb = load_workbook(XLSX, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"❌ シート「{SHEET}」が見つかりません（シート名を変えないでください）")
        sys.exit(1)
    ws = wb[SHEET]

    header_row = [text(c.value) for c in ws[1]]
    col = {}
    for h in HEADERS:
        if h not in header_row:
            print(f"❌ 1行目に見出し「{h}」がありません（見出しは変えないでください）")
            sys.exit(1)
        col[h] = header_row.index(h)

    for h in OPTIONAL_HEADERS:
        if h in header_row:
            col[h] = header_row.index(h)

    errors, warnings, tools, seen = [], [], [], {}

    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {h: text(row[i]) if i < len(row) else "" for h, i in col.items()}
        for h in OPTIONAL_HEADERS:
            values.setdefault(h, "")
        if not any(values.values()):
            continue  # 空の行は無視

        where = f"{r}行目（{values['AI名'] or '名前なし'}）"
        for h in REQUIRED:
            if not values[h]:
                errors.append(f"{where}：「{h}」が空です")

        name = values["AI名"]
        if name:
            if name in seen:
                errors.append(f"{where}：AI名が {seen[name]}行目 と重複しています")
            seen[name] = r

        url = values["URL"]
        if url and not re.match(r"^https?://", url):
            errors.append(f"{where}：URLは http:// か https:// で始めてください")

        pricing = values["料金タイプ"]
        if pricing and pricing not in PRICING:
            errors.append(f"{where}：料金タイプは「完全無料」「無料枠あり」「トライアルのみ」のどれかにしてください（今：{pricing}）")

        jp = normalize_japanese(values["日本語対応"])
        if jp and jp not in {"◯", "△", "×"}:
            errors.append(f"{where}：日本語対応は ◯ △ × のどれかにしてください（今：{values['日本語対応']}）")

        platform = values["対応端末"].replace("+", "＋").replace(" ", "").replace("　", "")
        if platform and platform not in PLATFORM:
            errors.append(f"{where}：対応端末は「PC＋スマホ」「PCのみ」「スマホのみ」のどれかにしてください（今：{values['対応端末']}）")

        cats = split_list(values["カテゴリ"], r"[、,，\n]")
        for c in cats:
            if c not in CATEGORIES:
                errors.append(f"{where}：カテゴリ「{c}」はサイトにありません（「カテゴリ一覧」シートの名前を使ってください）")

        rank = None
        if values["TOP10順位"]:
            try:
                rank = int(float(values["TOP10順位"]))
                if not 1 <= rank <= 10:
                    raise ValueError
            except ValueError:
                errors.append(f"{where}：TOP10順位は1〜10の数字にしてください")
                rank = None
        if rank and TOP10 not in cats:
            cats.insert(0, TOP10)
        if not rank and TOP10 in cats:
            warnings.append(f"{where}：カテゴリにTOP10がありますが、TOP10順位が空です")

        strengths = split_list(values["強み"], r"[\n／]")
        if strengths and len(strengths) != 4:
            warnings.append(f"{where}：強みが{len(strengths)}項目です（4項目がおすすめ）")

        tool = {
            "id": 0,
            "name": name,
            "url": url,
            "country": values["国"],
            "company": values["開発企業"],
            "category": cats,
            "strengths": strengths,
            "pricing": pricing,
            "freeLimit": values["無料利用範囲"] or DEFAULT_LIMIT.get(pricing, ""),
            "japanese": jp,
            "noSignup": to_bool(values["登録不要"]),
            "platform": PLATFORM.get(platform, "both"),
        }
        if values["AI名（英語）"]:
            tool["nameEn"] = values["AI名（英語）"]
        strengths_en = split_list(values["強み（英語）"], r"\n")
        if strengths_en:
            tool["strengthsEn"] = strengths_en
        elif "強み（英語）" in col:
            warnings.append(f"{where}：強み（英語）が空です（英語表示では日本語の強みが出ます）")
        if values["無料利用範囲（英語）"]:
            tool["freeLimitEn"] = values["無料利用範囲（英語）"]
        if rank:
            tool["rank"] = rank
        tools.append(tool)

    ranks = [t["rank"] for t in tools if "rank" in t]
    dup = {x for x in ranks if ranks.count(x) > 1}
    if dup:
        errors.append(f"TOP10順位が重複しています：{sorted(dup)}")

    for w in warnings:
        print(f"⚠️ {w}")

    if errors:
        print("\n❌ 入力ミスが見つかったため、JSONは更新しませんでした：")
        for e in errors:
            print(f"  ・{e}")
        sys.exit(1)

    for i, t in enumerate(tools, start=1):
        t["id"] = i

    jst = timezone(timedelta(hours=9))
    data = {"lastUpdated": datetime.now(jst).strftime("%Y-%m-%d"), "tools": tools}
    JSON_OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"✅ {len(tools)}件のAIを {JSON_OUT.name} に書き出しました")


if __name__ == "__main__":
    main()
